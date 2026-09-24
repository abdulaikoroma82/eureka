"""Tests for the NTD domain rule pack (``knowledge/packs/ntd.yaml``).

Positive cases prove each NTD label gets the pack's type and bound; the
``false_positive`` cases run the same helper over near-miss labels (reasons,
recall periods, yes/no drug questions) that must NOT be given a tablet or
parasitology bound, with the positive cases acting as their control."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from xlsform_studio.app.main import main
from xlsform_studio.app.simulator import Interview
from xlsform_studio.app.workflow import Workflow
from xlsform_studio.engine.knowledge_base import KnowledgeBase
from xlsform_studio.models import Question, Questionnaire


def _compile(labels, packs=("ntd",)):
    """Run labels (str, or (label, choices)) through the deterministic
    workflow with *packs* and return the result."""
    questions = []
    for i, item in enumerate(labels):
        label, choices = item if isinstance(item, tuple) else (item, [])
        questions.append(Question(name=f"q{i}", raw_label=label, label=label,
                                  raw_choices=list(choices)))
    kb = KnowledgeBase.load(packs=list(packs))
    return Workflow(knowledge=kb).run(Questionnaire(questions=questions),
                                      form_title="NTD", form_id="ntd",
                                      write_outputs=False)


POSITIVE = [
    # label, expected type, expected constraint
    ("How many praziquantel tablets did you swallow?", "decimal",
     ". >= 0 and . <= 5 and (. * 2) mod 1 = 0"),
    ("Number of PZQ tablets given", "decimal",
     ". >= 0 and . <= 5 and (. * 2) mod 1 = 0"),
    ("Number of ivermectin tablets given", "integer", ". >= 0 and . <= 4"),
    ("How many Mectizan tablets were swallowed?", "integer",
     ". >= 0 and . <= 4"),
    ("How many albendazole tablets did you swallow?", "decimal",
     ". >= 0 and . <= 1 and (. * 2) mod 1 = 0"),
    ("Number of azithromycin tablets given", "integer", ". >= 0 and . <= 4"),
    ("S. haematobium eggs per 10 ml urine", "integer",
     ". >= 0 and . <= 10000"),
    ("Ascaris eggs per gram (EPG)", "integer", ". >= 0 and . <= 600000"),
    ("Hookworm eggs counted on slide 1", "integer", ". >= 0 and . <= 25000"),
    ("Microfilariae per ml (night blood)", "decimal",
     ". >= 0 and . <= 50000"),
    ("Skin snip density (mf/mg)", "decimal", ". >= 0 and . <= 1000"),
    ("Number of lashes touching the eyeball (right eye)", "integer",
     ". >= 0 and . <= 200"),
    ("Height (cm) on the dose pole", "decimal", ". >= 30 and . <= 250"),
]


def test_ntd_pack_is_available_and_loads():
    assert "ntd" in KnowledgeBase.available_packs()
    kb = KnowledgeBase.load(packs=["ntd"])
    assert kb.packs == ["ntd"]


@pytest.mark.parametrize("label,xtype,constraint", POSITIVE,
                         ids=[p[0][:30] for p in POSITIVE])
def test_ntd_end_to_end_type_and_bound(label, xtype, constraint):
    q = _compile([label]).questionnaire.questions[0]
    assert q.xlsform_type == xtype
    assert q.constraint == constraint
    assert q.constraint_message


def test_ntd_end_to_end_form_is_valid_with_no_warnings():
    """Every pack bound together compiles, passes pyxform, and uses only
    expression syntax the tool can check (no 'unsupported' findings)."""
    result = _compile([p[0] for p in POSITIVE])
    assert result.is_valid
    assert any("Deep validation passed" in f.message
               for f in result.report.findings)
    assert [f.message for f in result.report.findings
            if f.level != "info"] == []


@pytest.mark.parametrize("value,ok", [
    ("0", True), ("0.5", True), ("2.5", True), ("5", True),
    ("2.3", False), ("5.5", False), ("-1", False)])
def test_ntd_praziquantel_half_tablet_bound_simulated(value, ok):
    """The half-tablet step rule behaves as intended when answered."""
    qn = _compile(["How many praziquantel tablets did you swallow?"])
    assert Interview(qn.questionnaire).submit(value).ok is ok


def test_ntd_neutral_bounds_unchanged_without_pack():
    q = _compile(["Number of ivermectin tablets given"],
                 packs=()).questionnaire.questions[0]
    assert q.constraint == ". >= 0"      # neutral "number of" rule only


def test_ntd_stacks_with_nutrition_pack():
    result = _compile(["MUAC", "How many praziquantel tablets did you swallow?"],
                      packs=("nutrition", "ntd"))
    constraints = [q.constraint for q in result.questionnaire.questions]
    assert constraints[0] == ". >= 60 and . <= 400"          # nutrition
    assert constraints[1].startswith(". >= 0 and . <= 5")    # ntd
    assert result.is_valid


NEAR_MISSES = [
    # label, choices, expected type, expected constraint
    ("Reason for not taking praziquantel", [], "text", ""),
    ("How many days ago did you take ivermectin?", [], "integer", ". >= 0"),
    ("Did you swallow the albendazole tablet?", ["Yes", "No"],
     "select_one yes_no", ""),
    ("Number of children positive for S. haematobium", [], "integer",
     ". >= 0"),
    ("Was the dose pole used?", ["Yes", "No"], "select_one yes_no", ""),
    ("Was a skin snip taken?", ["Yes", "No"], "select_one yes_no", ""),
]


@pytest.mark.parametrize("label,choices,xtype,constraint", NEAR_MISSES,
                         ids=[n[0][:30] for n in NEAR_MISSES])
def test_ntd_false_positive_labels_not_bounded(label, choices, xtype,
                                               constraint):
    q = _compile([(label, choices)]).questionnaire.questions[0]
    assert q.xlsform_type == xtype
    assert q.constraint == constraint


def test_ntd_cli_packs_flag_writes_bounded_xlsform(tmp_path: Path):
    source = tmp_path / "mda_coverage.csv"
    source.write_text(
        "question,type,choices,required,section,logic\n"
        "Did you swallow the MDA drugs?,,Yes|No,yes,MDA,\n"
        "How many praziquantel tablets did you swallow?,,,yes,MDA,\n"
        "Number of ivermectin tablets given,,,yes,MDA,\n",
        encoding="utf-8")
    out = tmp_path / "out"
    assert main([str(source), "--packs", "ntd", "--output", str(out),
                 "--quiet", "--no-path-analysis"]) == 0

    # The package holds several workbooks; the XLSForm is the one with a
    # 'survey' sheet.
    books = [openpyxl.load_workbook(p) for p in out.rglob("*.xlsx")]
    sheet = next(b["survey"] for b in books if "survey" in b.sheetnames)
    header = [c.value for c in sheet[1]]
    rows = [dict(zip(header, (c.value for c in row)))
            for row in sheet.iter_rows(min_row=2)]
    by_label = {r.get("label") or r.get("label::English (en)"): r
                for r in rows}
    pzq = next(r for lbl, r in by_label.items()
               if lbl and "praziquantel" in lbl)
    ivm = next(r for lbl, r in by_label.items()
               if lbl and "ivermectin" in lbl)
    assert pzq["constraint"] == ". >= 0 and . <= 5 and (. * 2) mod 1 = 0"
    assert ivm["constraint"] == ". >= 0 and . <= 4"


def test_ntd_cli_unknown_pack_is_rejected(tmp_path: Path, capsys):
    """A typo in --packs fails loudly and lists ntd among the options."""
    source = tmp_path / "f.csv"
    source.write_text("question\nAge\n", encoding="utf-8")
    assert main([str(source), "--packs", "ntds", "--output",
                 str(tmp_path / "o"), "--quiet"]) != 0
    err = capsys.readouterr().err
    assert "Unknown domain rule pack 'ntds'" in err
    assert "ntd" in err.split("Available:")[1]
